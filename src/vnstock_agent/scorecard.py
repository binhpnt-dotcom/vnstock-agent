"""FO Market Scorecard automation (Polo cycle/allocation model).

This module drives the "Market_Scorecard_Polo" Excel workbook: it fetches
live VN market data via vnstock, recomputes the scorecard's formulas in
pure Python (so cycle phase / allocation numbers are available without
opening Excel), writes the fetched inputs back into the workbook's raw
input cells (leaving every formula untouched so the file still recalculates
correctly if opened in Excel/Sheets/LibreOffice), and renders an HTML
dashboard.

The formula replication below was verified cell-by-cell against the cached
values embedded in the original workbook (P/E z-score, group scores,
overall score, stance, suggested equity weight, allocation table, cycle
phase, etc. all matched to float precision). Two latent issues were found
in the original sheet during that verification and are surfaced here as
`ScorecardResult.warnings` rather than silently "fixed":

1. Cell E8 holds the text "Macro Regime" (a stray table header) instead of
   a numeric weight, even though the Overall Market Score formula
   (`=SUMPRODUCT(B53:B57,E4:E8)`) multiplies it against the CyclePsych
   group score. A non-numeric operand contributes 0, so the CyclePsych
   group is silently excluded from the Overall Market Score even though it
   still drives Sentiment Regime / Cycle Phase.
2. Row 27's Group label is `"Flow "` (trailing space) instead of `"Flow"`,
   so `SUMIF(A17:A49,"Flow",...)` excludes it — "Liquidity vs 6M avg" does
   not contribute to the Flow group score.
"""

from __future__ import annotations

import concurrent.futures
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

import openpyxl
import pandas as pd

from vnstock_agent import core
from vnstock_agent.config import DEFAULT_SOURCE

SHEET_NAME = "Market_Scorecard_Polo"
INDICATOR_ROWS = range(17, 50)  # rows 17-49 inclusive

GROUPS = ["Value", "Technical", "Flow", "Macro", "CyclePsych"]

# Rows whose CurrentValue (col C) formula we recompute ourselves from base
# inputs instead of treating as a manual entry.
_DERIVED_ROWS = {17, 18, 19, 20, 21, 25, 26, 27, 33, 39, 44}
# Rows whose CurrentValue we overwrite with live market data.
_AUTO_ROWS = {22, 23, 24, 47}
_SAFE_ARITHMETIC_RE = re.compile(r"^[0-9.\s+\-*/()]+$")

# Cell locations for the base/auto inputs.
CELL_VNINDEX_SPOT = "B4"
CELL_LIQ_20D = "B5"
CELL_FX_BASELINE = "B6"
CELL_FX_TODAY = "B7"
CELL_LIQ_6M_AVG = "I9"
CELL_MA50 = "H10"
CELL_MA200 = "H11"
CELL_MA50_20D_AGO = "I10"
CELL_LAST_UPDATED = "B2"
CELL_AS_OF_DATE = "B3"


def _num(value, default=0.0) -> float:
    if value is None:
        return default
    if hasattr(value, "text"):  # ArrayFormula
        value = value.text
    if isinstance(value, str):
        s = value.strip()
        if s.startswith("="):
            s = s[1:]
        if _SAFE_ARITHMETIC_RE.match(s):
            try:
                return float(eval(s, {"__builtins__": {}}, {}))  # noqa: S307 whitelisted chars only
            except Exception:
                return default
        try:
            return float(s)
        except ValueError:
            return default
    return float(value)


def _clamp(v: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, v))


@dataclass
class IndicatorRow:
    row: int
    group: str
    factor: str
    current: float
    weight: float
    good: float
    bad: float
    lower_is_better: bool
    source: str  # "auto" | "derived" | "manual"
    score10: float = 0.0
    score100: float = 0.0


@dataclass
class BaseInputs:
    vnindex_spot: float
    liq_20d: float
    fx_baseline: float
    fx_today: float
    liq_6m_avg: float
    ma50: float
    ma200: float
    ma50_20d_ago: float
    pe_current: float
    pe_avg5y: float
    pb_current: float
    pb_avg5y: float
    growth_avg_pct: float
    roe_pct: float
    tpcp_10y_pct: float
    group_weight: dict  # {"Value": .., "Technical": .., "Flow": .., "Macro": .., "CyclePsych": raw}


@dataclass
class AutoData:
    """Live-fetched values. None means the fetch failed / was skipped."""

    vnindex_spot: Optional[float] = None
    liq_20d: Optional[float] = None
    liq_6m_avg: Optional[float] = None
    ma50: Optional[float] = None
    ma200: Optional[float] = None
    ma50_20d_ago: Optional[float] = None
    rsi14: Optional[float] = None
    pct_above_ma50: Optional[float] = None
    pct_above_ma200: Optional[float] = None
    breadth_sample_size: Optional[int] = None
    btc_roc_30d: Optional[float] = None
    fx_today: Optional[float] = None
    errors: dict = field(default_factory=dict)


@dataclass
class ScorecardResult:
    as_of: str
    rows: list  # list[IndicatorRow]
    group_scores: dict
    overall_score: float
    stance: str
    macro_regime: str
    sentiment_regime: str
    cycle_phase: str
    suggested_equity_weight: float
    leverage_note: str
    allocation: list  # list[dict]
    narrative: str
    liquidity_band: str
    warnings: list


# --- Reading the workbook -------------------------------------------------


def _read_row_meta(ws, row: int) -> dict:
    return dict(
        row=row,
        group=ws[f"A{row}"].value,
        factor=ws[f"B{row}"].value,
        weight=_num(ws[f"F{row}"].value),
        good=_num(ws[f"H{row}"].value),
        bad=_num(ws[f"I{row}"].value),
        lower_is_better=(ws[f"J{row}"].value == 1),
    )


def read_base_inputs(ws) -> BaseInputs:
    return BaseInputs(
        vnindex_spot=_num(ws[CELL_VNINDEX_SPOT].value),
        liq_20d=_num(ws[CELL_LIQ_20D].value),
        fx_baseline=_num(ws[CELL_FX_BASELINE].value),
        fx_today=_num(ws[CELL_FX_TODAY].value),
        liq_6m_avg=_num(ws[CELL_LIQ_6M_AVG].value),
        ma50=_num(ws[CELL_MA50].value),
        ma200=_num(ws[CELL_MA200].value),
        ma50_20d_ago=_num(ws[CELL_MA50_20D_AGO].value),
        pe_current=_num(ws["C14"].value),
        pe_avg5y=_num(ws["D14"].value),
        pb_current=_num(ws["C15"].value),
        pb_avg5y=_num(ws["D15"].value),
        growth_avg_pct=_num(ws["I14"].value),
        roe_pct=_num(ws["I15"].value),
        tpcp_10y_pct=_num(ws["H12"].value),
        group_weight={
            "Value": _num(ws["E4"].value),
            "Technical": _num(ws["E5"].value),
            "Flow": _num(ws["E6"].value),
            "Macro": _num(ws["E7"].value),
            "CyclePsych": ws["E8"].value,
        },
    )


# --- Fetching live data via vnstock ---------------------------------------


def _history_df(symbol: str, days: int, source: str) -> pd.DataFrame:
    from datetime import timedelta

    end = datetime.now().strftime("%Y-%m-%d")
    start = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d")
    records = core.stock_history(symbol, start, end, "1D", source)
    df = pd.DataFrame(records)
    if df.empty or "close" not in df.columns:
        raise ValueError(f"no usable history returned for {symbol}")
    time_col = "time" if "time" in df.columns else df.columns[0]
    df = df.sort_values(time_col).reset_index(drop=True)
    df["close"] = df["close"].astype(float)
    return df


def _rsi14(closes: pd.Series, period: int = 14) -> float:
    """Cutler's RSI (SMA-smoothed gains/losses) over the last `period` bars."""
    delta = closes.diff().dropna()
    gains = delta.clip(lower=0)
    losses = -delta.clip(upper=0)
    avg_gain = gains.rolling(period).mean().iloc[-1]
    avg_loss = losses.rolling(period).mean().iloc[-1]
    if avg_loss == 0:
        return 100.0
    rs = avg_gain / avg_loss
    return 100 - 100 / (1 + rs)


def _turnover_value_bn(df: pd.DataFrame) -> Optional[pd.Series]:
    """Daily traded value in VND billion, if the source provides it."""
    for col in ("value", "match_value", "total_value", "totalValue"):
        if col in df.columns:
            return df[col].astype(float) / 1e9
    return None


def _fetch_index_technicals(source: str, errors: dict) -> dict:
    out = {}
    try:
        df = _history_df("VNINDEX", days=420, source=source)
        closes = df["close"]
        out["vnindex_spot"] = float(closes.iloc[-1])
        out["ma50"] = float(closes.rolling(50).mean().iloc[-1])
        out["ma200"] = float(closes.rolling(200).mean().iloc[-1])
        ma50_series = closes.rolling(50).mean()
        if len(ma50_series.dropna()) > 20:
            out["ma50_20d_ago"] = float(ma50_series.iloc[-21])
        out["rsi14"] = float(_rsi14(closes))

        turnover = _turnover_value_bn(df)
        if turnover is not None:
            out["liq_20d"] = float(turnover.tail(20).mean())
            if len(turnover) >= 40:
                out["liq_6m_avg"] = float(turnover.tail(126).mean())
        else:
            errors["liquidity"] = (
                "History source has no traded-value column (checked value/match_value/"
                "total_value); B5/I9 left as manual input."
            )
    except Exception as e:
        errors["index_technicals"] = str(e)
    return out


def _fetch_breadth(source: str, errors: dict) -> dict:
    """% of VN30 members trading above their own MA50 / MA200, as a market-breadth proxy."""
    out = {}
    try:
        symbols_records = core.listing_symbols_by_group("VN30", source)
        symbols = [r.get("symbol") or r.get("ticker") for r in symbols_records if r]
        symbols = [s for s in symbols if s]
        if not symbols:
            raise ValueError("empty VN30 constituent list")

        def _above_ma(sym):
            df = _history_df(sym, days=420, source=source)
            closes = df["close"]
            last = closes.iloc[-1]
            ma50 = closes.rolling(50).mean().iloc[-1]
            ma200 = closes.rolling(200).mean().iloc[-1]
            return (last > ma50) if pd.notna(ma50) else None, (last > ma200) if pd.notna(ma200) else None

        above50, above200 = [], []
        with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
            futures = {pool.submit(_above_ma, s): s for s in symbols}
            for fut in concurrent.futures.as_completed(futures):
                try:
                    a50, a200 = fut.result()
                    if a50 is not None:
                        above50.append(a50)
                    if a200 is not None:
                        above200.append(a200)
                except Exception:
                    continue

        if len(above50) < len(symbols) // 2:
            raise ValueError(f"too few VN30 symbols resolved ({len(above50)}/{len(symbols)})")

        out["pct_above_ma50"] = sum(above50) / len(above50)
        out["pct_above_ma200"] = sum(above200) / len(above200)
        out["breadth_sample_size"] = len(above50)
    except Exception as e:
        errors["breadth"] = f"{e} (C23/C24 left as manual input; this is a VN30 proxy, not full-market breadth)"
    return out


def _fetch_fx(errors: dict) -> Optional[float]:
    try:
        records = core.fx_history("USDVND", None, None, "1D")
        df = pd.DataFrame(records)
        if df.empty or "error" in df.columns:
            raise ValueError(str(records[:1]))
        close_col = "close" if "close" in df.columns else df.columns[-1]
        time_col = "time" if "time" in df.columns else df.columns[0]
        df = df.sort_values(time_col)
        return float(df[close_col].iloc[-1])
    except Exception as e:
        errors["fx"] = str(e)
        return None


def _fetch_btc_roc_30d(errors: dict) -> Optional[float]:
    try:
        records = core.crypto_history("BTC", None, None, "1D")
        df = pd.DataFrame(records)
        if df.empty or "close" not in df.columns or len(df) < 31:
            raise ValueError("insufficient BTC history")
        time_col = "time" if "time" in df.columns else df.columns[0]
        df = df.sort_values(time_col).reset_index(drop=True)
        closes = df["close"].astype(float)
        return float((closes.iloc[-1] - closes.iloc[-31]) / closes.iloc[-31])
    except Exception as e:
        errors["btc"] = str(e)
        return None


def fetch_auto_data(source: str = DEFAULT_SOURCE, include_breadth: bool = True) -> AutoData:
    """Fetch everything this module can reliably pull from vnstock.

    Any single failure is caught and reported in `AutoData.errors`; the
    corresponding field stays None so the caller keeps the workbook's
    existing manual value instead of overwriting it with bad data.
    """
    errors: dict = {}
    idx = _fetch_index_technicals(source, errors)
    breadth = _fetch_breadth(source, errors) if include_breadth else {}
    fx_today = _fetch_fx(errors)
    btc_roc = _fetch_btc_roc_30d(errors)

    return AutoData(
        vnindex_spot=idx.get("vnindex_spot"),
        liq_20d=idx.get("liq_20d"),
        liq_6m_avg=idx.get("liq_6m_avg"),
        ma50=idx.get("ma50"),
        ma200=idx.get("ma200"),
        ma50_20d_ago=idx.get("ma50_20d_ago"),
        rsi14=idx.get("rsi14"),
        pct_above_ma50=breadth.get("pct_above_ma50"),
        pct_above_ma200=breadth.get("pct_above_ma200"),
        breadth_sample_size=breadth.get("breadth_sample_size"),
        btc_roc_30d=btc_roc,
        fx_today=fx_today,
        errors=errors,
    )


# --- Scoring engine ---------------------------------------------------------


def _derived_current(row: int, base: BaseInputs, manual: dict) -> float:
    """Recompute the CurrentValue for rows whose Excel formula depends on
    other (now auto-updated) cells, mirroring the sheet's own formulas."""
    if row == 17:
        return base.pe_current - base.pe_avg5y
    if row == 18:
        return base.pb_current - base.pb_avg5y
    if row == 19:
        return base.pe_current / base.growth_avg_pct / 100 if base.growth_avg_pct else 0.0
    if row == 20:
        return base.roe_pct
    if row == 21:
        return 11 + base.tpcp_10y_pct
    if row == 25:
        return base.vnindex_spot / base.ma50 if base.ma50 else 0.0
    if row == 26:
        return (base.ma50 - base.ma50_20d_ago) / base.ma50_20d_ago if base.ma50_20d_ago else 0.0
    if row == 27:
        return base.liq_20d / base.liq_6m_avg if base.liq_6m_avg else 0.0
    if row == 33:
        return base.liq_20d
    if row == 39:
        return base.fx_today / base.fx_baseline - 1 if base.fx_baseline else 0.0
    if row == 44:
        return (base.pe_current - base.pe_avg5y) / 2.3
    raise ValueError(f"row {row} is not a derived-formula row")


def _score10(meta: dict, current: float) -> float:
    if meta["row"] == 33:
        # Band scoring (Excel: IF(C<=15000,10,IF(C<=25000,7,IF(C<=40000,4,2))))
        if current <= 15000:
            return 10.0
        if current <= 25000:
            return 7.0
        if current <= 40000:
            return 4.0
        return 2.0
    good, bad, lower = meta["good"], meta["bad"], meta["lower_is_better"]
    denom = (bad - good) if lower else (good - bad)
    if denom == 0:
        return 0.0
    raw = (bad - current) / denom * 10 if lower else (current - bad) / denom * 10
    return _clamp(raw, 0.0, 10.0)


def build_rows(ws, base: BaseInputs, auto: AutoData) -> list:
    auto_map = {
        22: auto.rsi14,
        23: auto.pct_above_ma50,
        24: auto.pct_above_ma200,
        47: auto.btc_roc_30d,
    }
    rows = []
    for r in INDICATOR_ROWS:
        meta = _read_row_meta(ws, r)
        if r in _DERIVED_ROWS:
            current = _derived_current(r, base, {})
            source_tag = "derived"
        elif r in _AUTO_ROWS and auto_map.get(r) is not None:
            current = auto_map[r]
            source_tag = "auto"
        else:
            current = _num(ws[f"C{r}"].value)
            source_tag = "manual" if r not in _AUTO_ROWS else "manual (auto-fetch unavailable)"
        score10 = _score10(meta, current)
        rows.append(
            IndicatorRow(
                row=r,
                group=meta["group"],
                factor=meta["factor"],
                current=current,
                weight=meta["weight"],
                good=meta["good"],
                bad=meta["bad"],
                lower_is_better=meta["lower_is_better"],
                source=source_tag,
                score10=score10,
                score100=score10 * 10,
            )
        )
    return rows


def _group_score(rows: list, group: str) -> float:
    items = [r for r in rows if r.group == group]
    wsum = sum(r.weight for r in items)
    if wsum == 0:
        return 0.0
    return sum(r.score100 * r.weight for r in items) / wsum


_LIQUIDITY_BANDS = [(15000, "10-15k"), (25000, "15-25k"), (40000, "25-40k"), (float("inf"), ">40k")]
_CYCLE_EQUITY_FACTOR = {
    "Fear / Capitulation": 1.4,
    "Early Recovery": 1.3,
    "Early Accumulation": 1.2,
    "Mid-cycle / Neutral": 1.0,
    "Late Expansion": 0.85,
    "Downturn / Repricing": 0.8,
    "FOMO / Blow-off": 0.7,
}


def _liquidity_band(liq_20d: float) -> str:
    for threshold, label in _LIQUIDITY_BANDS:
        if liq_20d <= threshold:
            return label
    return ">40k"


def _macro_regime(rows: list) -> str:
    macro_score10 = [r.score10 for r in rows if r.row in (35, 36, 37, 38)]
    if not macro_score10:
        return ""
    avg = sum(macro_score10) / len(macro_score10)
    if avg < 4:
        return "Tight"
    if avg < 6:
        return "Neutral"
    if avg <= 10:
        return "Loose"
    return ""


def _sentiment_regime(cyclepsych_score: float) -> str:
    if cyclepsych_score < 40:
        return "Fear"
    if cyclepsych_score <= 70:
        return "Neutral"
    return "Greed"


def _cycle_phase(liq_20d: float, sentiment: str, macro: str, adr: float) -> str:
    if liq_20d < 15000 and sentiment == "Fear" and macro == "Tight":
        return "Fear / Capitulation"
    if liq_20d < 25000 and sentiment == "Neutral":
        return "Early Accumulation"
    if liq_20d < 35000 and sentiment == "Neutral" and adr >= 0.9:
        return "Early Momentum"
    if liq_20d < 40000 and sentiment == "Greed" and macro != "Tight":
        return "Late Expansion"
    if liq_20d >= 40000 and sentiment == "Greed":
        return "FOMO / Blow-off"
    if liq_20d >= 25000 and sentiment != "Greed" and adr < 0.9:
        return "Downturn / Repricing"
    return "Mid-cycle / Neutral"


def compute(ws, base: BaseInputs, auto: AutoData) -> ScorecardResult:
    warns = list(auto.errors.values())

    rows = build_rows(ws, base, auto)
    group_scores = {g: _group_score(rows, g) for g in GROUPS}

    cyclepsych_weight_raw = base.group_weight["CyclePsych"]
    cyclepsych_weight = cyclepsych_weight_raw if isinstance(cyclepsych_weight_raw, (int, float)) else 0.0
    if not isinstance(cyclepsych_weight_raw, (int, float)):
        warns.append(
            f"Cell E8 holds {cyclepsych_weight_raw!r} (not a number). The Overall Market Score "
            "formula multiplies it by the CyclePsych group score, so CyclePsych is currently "
            "contributing 0% to the Overall Score even though it still drives Sentiment Regime / "
            "Cycle Phase. Put a numeric weight in E8 if that group should count toward the score."
        )
    flow_row27 = next((r for r in rows if r.row == 27), None)
    if flow_row27 and flow_row27.group != "Flow":
        warns.append(
            f"Row 27's Group cell (A27) is {flow_row27.group!r} instead of \"Flow\", so "
            "'Liquidity vs 6M avg (ratio)' is excluded from the Flow group score (Excel's "
            "SUMIF only matches exact text)."
        )

    overall_score = (
        group_scores["Value"] * base.group_weight["Value"]
        + group_scores["Technical"] * base.group_weight["Technical"]
        + group_scores["Flow"] * base.group_weight["Flow"]
        + group_scores["Macro"] * base.group_weight["Macro"]
        + group_scores["CyclePsych"] * cyclepsych_weight
    )

    if overall_score >= 80:
        stance = "Rủi ro (giảm beta)"
    elif overall_score >= 60:
        stance = "Trung tính (giữ barbell)"
    else:
        stance = "Tích lũy (mở vị thế)"

    macro_regime = _macro_regime(rows)
    sentiment_regime = _sentiment_regime(group_scores["CyclePsych"])
    adr_row = next((r for r in rows if r.row == 28), None)
    adr_value = adr_row.current if adr_row else 0.0
    cycle_phase = _cycle_phase(base.liq_20d, sentiment_regime, macro_regime, adr_value)

    factor = _CYCLE_EQUITY_FACTOR.get(cycle_phase, 1.0)
    suggested_equity_weight = _clamp((overall_score / 100) * factor, 0.0, 1.5) * 1.2
    leverage_note = "zero" if suggested_equity_weight <= 1 else f"{suggested_equity_weight - 1:.2%} margin"

    if overall_score < 60:
        core_pct, growth_pct = 0.7, 0.2
    elif overall_score > 80:
        core_pct, growth_pct = 0.5, 0.3
    else:
        core_pct, growth_pct = 0.6, 0.25
    lynch_pct = max(0.0, 1 - core_pct - growth_pct)

    allocation = [
        {
            "layer": "Core Compounders (C)",
            "pct_equity": core_pct,
            "pct_nav": core_pct * suggested_equity_weight,
            "notes": "FPT, REE, ACB, SIP, VN30 quality",
        },
        {
            "layer": "Growth Catalysts (G)",
            "pct_equity": growth_pct,
            "pct_nav": growth_pct * suggested_equity_weight,
            "notes": "Banks mid/high beta, KCN, oil&gas, logistics, retails",
        },
        {
            "layer": "Lynch/Fisher Layer (L)",
            "pct_equity": lynch_pct,
            "pct_nav": lynch_pct * suggested_equity_weight,
            "notes": "Niche small/mid, only when liquidity not overheated",
        },
        {
            "layer": "Cash / Defensive",
            "pct_equity": None,
            "pct_nav": 1 - suggested_equity_weight,
            "notes": "Used to de-risk in late-cycle / FOMO phase",
        },
    ]

    if overall_score >= 80:
        narrative = "Late-cycle / FOMO – Giảm beta, tăng C-layer, nâng tiền mặt."
    elif overall_score >= 60:
        narrative = "Mid-cycle – Giữ barbell cân bằng, xoay G-layer có chọn lọc."
    else:
        narrative = "Early recovery / Fear zone – Tăng giải ngân G-layer, ưu tiên value & quality."

    return ScorecardResult(
        as_of=datetime.now().strftime("%Y-%m-%d"),
        rows=rows,
        group_scores=group_scores,
        overall_score=overall_score,
        stance=stance,
        macro_regime=macro_regime,
        sentiment_regime=sentiment_regime,
        cycle_phase=cycle_phase,
        suggested_equity_weight=suggested_equity_weight,
        leverage_note=leverage_note,
        allocation=allocation,
        narrative=narrative,
        liquidity_band=_liquidity_band(base.liq_20d),
        warnings=warns,
    )


# --- Workbook I/O ------------------------------------------------------------


_AUTO_FIELD_CELLS = {
    "vnindex_spot": CELL_VNINDEX_SPOT,
    "liq_20d": CELL_LIQ_20D,
    "liq_6m_avg": CELL_LIQ_6M_AVG,
    "ma50": CELL_MA50,
    "ma200": CELL_MA200,
    "ma50_20d_ago": CELL_MA50_20D_AGO,
    "rsi14": "C22",
    "pct_above_ma50": "C23",
    "pct_above_ma200": "C24",
    "btc_roc_30d": "C47",
    "fx_today": CELL_FX_TODAY,
}


def apply_auto_data(ws, auto: AutoData) -> list:
    """Write fetched values into the workbook's raw input cells (formulas
    elsewhere are left untouched). Returns the list of fields actually applied."""
    applied = []
    for field_name, cell in _AUTO_FIELD_CELLS.items():
        value = getattr(auto, field_name)
        if value is not None:
            ws[cell] = value
            applied.append(field_name)
    ws[CELL_LAST_UPDATED] = datetime.now()
    ws[CELL_AS_OF_DATE] = datetime.now().strftime("%Y-%m-%d")
    return applied


def load_workbook_sheet(path: str):
    wb = openpyxl.load_workbook(path, data_only=False)
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.active
    return wb, ws


def update_workbook(
    path: str,
    out_path: Optional[str] = None,
    source: str = DEFAULT_SOURCE,
    include_breadth: bool = True,
) -> tuple:
    """Fetch live data, recompute the scorecard, write it back to the workbook.

    Returns (ScorecardResult, AutoData, applied_fields).
    """
    wb, ws = load_workbook_sheet(path)
    auto = fetch_auto_data(source=source, include_breadth=include_breadth)
    applied = apply_auto_data(ws, auto)

    # Re-read base inputs after writing so derived formulas use fresh values.
    base = read_base_inputs(ws)
    result = compute(ws, base, auto)

    wb.save(out_path or path)
    return result, auto, applied


# --- HTML dashboard rendering ------------------------------------------------

# Fixed-order categorical slots (see dataviz palette) — assigned by identity,
# never re-cycled based on value/rank.
_GROUP_COLORS = {
    "Value": "#2a78d6",  # blue
    "Technical": "#1baf7a",  # aqua
    "Flow": "#eda100",  # yellow
    "Macro": "#008300",  # green
    "CyclePsych": "#4a3aa7",  # violet
}
_LAYER_COLORS = {
    "Core Compounders (C)": "#2a78d6",
    "Growth Catalysts (G)": "#1baf7a",
    "Lynch/Fisher Layer (L)": "#eda100",
    "Cash / Defensive": "#898781",
}
_STANCE_STATUS = {
    "Tích lũy (mở vị thế)": ("good", "#0ca30c"),
    "Trung tính (giữ barbell)": ("warning", "#fab219"),
    "Rủi ro (giảm beta)": ("critical", "#d03b3b"),
}


def _esc(s) -> str:
    import html

    return html.escape(str(s))


def _fmt_pct(v: Optional[float], digits: int = 1) -> str:
    if v is None:
        return "—"
    return f"{v * 100:.{digits}f}%"


def _fmt_num(v: float, digits: int = 2) -> str:
    return f"{v:,.{digits}f}"


def render_dashboard_html(
    result: ScorecardResult,
    auto: AutoData,
    applied_fields: list,
    title: str = "FO Market Scorecard – Polo Investing",
) -> str:
    stance_key, stance_color = _STANCE_STATUS.get(result.stance, ("warning", "#fab219"))

    group_rows_html = ""
    for g in GROUPS:
        score = result.group_scores[g]
        color = _GROUP_COLORS[g]
        width = max(2, min(100, score))
        group_rows_html += f"""
        <div class="bar-row">
          <div class="bar-label"><span class="dot" style="background:{color}"></span>{_esc(g)}</div>
          <div class="bar-track">
            <div class="bar-fill" style="width:{width}%;background:{color}"></div>
          </div>
          <div class="bar-value">{score:.1f}</div>
        </div>"""

    alloc_segments = ""
    alloc_legend = ""
    for a in result.allocation:
        pct_nav = a["pct_nav"] or 0.0
        color = _LAYER_COLORS[a["layer"]]
        alloc_segments += (
            f'<div class="seg" style="width:{max(0, pct_nav) * 100:.2f}%;background:{color}" '
            f'title="{_esc(a["layer"])}: {_fmt_pct(pct_nav)} of NAV"></div>'
        )
        alloc_legend += f"""
        <div class="legend-item">
          <span class="dot" style="background:{color}"></span>
          <span class="legend-label">{_esc(a["layer"])}</span>
          <span class="legend-value">{_fmt_pct(pct_nav)} NAV{
            f' &middot; {_fmt_pct(a["pct_equity"])} equity' if a["pct_equity"] is not None else ''
        }</span>
        </div>"""

    auto_badge_map = {
        "vnindex_spot": "VNIndex spot (B4)",
        "liq_20d": "20D avg liquidity (B5)",
        "liq_6m_avg": "6M avg liquidity (I9)",
        "ma50": "MA50 (H10)",
        "ma200": "MA200 (H11)",
        "ma50_20d_ago": "MA50, 20d ago (I10)",
        "rsi14": "RSI 14D (C22)",
        "pct_above_ma50": "% > MA50 (C23, VN30 proxy)",
        "pct_above_ma200": "% > MA200 (C24, VN30 proxy)",
        "btc_roc_30d": "BTC 30D ROC (C47)",
        "fx_today": "USD/VND today (B7)",
    }
    auto_list_html = "".join(f"<li>{_esc(auto_badge_map.get(f, f))}</li>" for f in applied_fields)
    if not auto_list_html:
        auto_list_html = "<li>No live data was fetched for this render — showing formulas over the workbook's existing manual values.</li>"

    errors_html = "".join(f"<li>{_esc(msg)}</li>" for msg in auto.errors.values())
    warnings_html = "".join(f"<li>{_esc(w)}</li>" for w in result.warnings if w not in auto.errors.values())

    indicator_rows_html = ""
    for r in sorted(result.rows, key=lambda x: x.row):
        source_class = {"auto": "tag-auto", "derived": "tag-derived"}.get(r.source, "tag-manual")
        source_label = {"auto": "auto", "derived": "derived"}.get(r.source, "manual")
        indicator_rows_html += f"""
        <tr>
          <td><span class="dot" style="background:{_GROUP_COLORS.get(r.group.strip(), '#898781')}"></span>{_esc(r.group)}</td>
          <td>{_esc(r.factor)}</td>
          <td class="num">{_fmt_num(r.current, 4)}</td>
          <td class="num">{r.score10:.1f}</td>
          <td><span class="tag {source_class}">{source_label}</span></td>
        </tr>"""

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!doctype html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_esc(title)}</title>
<style>
  :root {{
    --surface-1: #fcfcfb;
    --page: #f9f9f7;
    --text-primary: #0b0b0b;
    --text-secondary: #52514e;
    --text-muted: #898781;
    --gridline: #e1e0d9;
    --border: rgba(11,11,11,0.10);
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0; padding: 32px 16px;
    background: var(--page);
    color: var(--text-primary);
    font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
  }}
  .wrap {{ max-width: 960px; margin: 0 auto; }}
  h1 {{ font-size: 20px; margin: 0 0 4px; }}
  .meta {{ color: var(--text-muted); font-size: 13px; margin-bottom: 24px; }}
  .card {{
    background: var(--surface-1);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 20px 24px;
    margin-bottom: 16px;
  }}
  .card h2 {{ font-size: 13px; text-transform: uppercase; letter-spacing: .04em; color: var(--text-secondary); margin: 0 0 16px; font-weight: 600; }}
  .hero-row {{ display: flex; gap: 16px; flex-wrap: wrap; }}
  .hero-tile {{ flex: 1; min-width: 180px; }}
  .hero-value {{ font-size: 48px; font-weight: 600; line-height: 1; }}
  .hero-label {{ color: var(--text-secondary); font-size: 13px; margin-top: 6px; }}
  .badge {{
    display: inline-block; padding: 4px 12px; border-radius: 999px;
    font-size: 13px; font-weight: 600; color: #fff; margin-top: 8px;
  }}
  .bar-row {{ display: flex; align-items: center; gap: 12px; margin-bottom: 12px; }}
  .bar-label {{ width: 130px; font-size: 13px; color: var(--text-secondary); display: flex; align-items: center; gap: 6px; flex-shrink: 0; }}
  .bar-track {{ flex: 1; height: 20px; background: var(--gridline); border-radius: 4px; overflow: hidden; }}
  .bar-fill {{ height: 100%; border-radius: 4px; }}
  .bar-value {{ width: 42px; text-align: right; font-size: 13px; font-variant-numeric: tabular-nums; color: var(--text-primary); }}
  .dot {{ display: inline-block; width: 8px; height: 8px; border-radius: 50%; flex-shrink: 0; }}
  .alloc-bar {{ display: flex; height: 24px; border-radius: 4px; overflow: hidden; gap: 2px; background: var(--surface-1); }}
  .seg {{ height: 100%; }}
  .legend {{ display: flex; flex-direction: column; gap: 8px; margin-top: 16px; }}
  .legend-item {{ display: flex; align-items: center; gap: 8px; font-size: 13px; }}
  .legend-label {{ color: var(--text-primary); min-width: 180px; }}
  .legend-value {{ color: var(--text-secondary); font-variant-numeric: tabular-nums; }}
  .narrative {{ font-size: 15px; line-height: 1.5; }}
  table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
  th, td {{ text-align: left; padding: 8px 10px; border-bottom: 1px solid var(--gridline); }}
  th {{ color: var(--text-muted); font-weight: 600; font-size: 12px; text-transform: uppercase; letter-spacing: .03em; }}
  td.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
  .tag {{ font-size: 11px; padding: 2px 8px; border-radius: 999px; font-weight: 600; }}
  .tag-auto {{ background: #cde2fb; color: #184f95; }}
  .tag-derived {{ background: #e1e0d9; color: #52514e; }}
  .tag-manual {{ background: #f0efec; color: #898781; }}
  ul {{ margin: 0; padding-left: 20px; font-size: 13px; color: var(--text-secondary); line-height: 1.6; }}
  .warn-list li {{ color: #9a5b00; }}
  .two-col {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; }}
  @media (max-width: 640px) {{ .two-col {{ grid-template-columns: 1fr; }} }}
</style>
</head>
<body>
<div class="wrap">
  <h1>{_esc(title)}</h1>
  <div class="meta">As of {_esc(result.as_of)} &middot; generated {generated_at}</div>

  <div class="card">
    <div class="hero-row">
      <div class="hero-tile">
        <div class="hero-value">{result.overall_score:.1f}</div>
        <div class="hero-label">Overall Market Score (0–100)</div>
        <div class="badge" style="background:{stance_color}">{_esc(result.stance)}</div>
      </div>
      <div class="hero-tile">
        <div class="hero-value" style="font-size:28px">{_esc(result.cycle_phase)}</div>
        <div class="hero-label">Cycle Phase &middot; {_esc(result.liquidity_band)} liquidity band</div>
        <div class="hero-label">Sentiment: {_esc(result.sentiment_regime)} &middot; Macro: {_esc(result.macro_regime)}</div>
      </div>
      <div class="hero-tile">
        <div class="hero-value" style="font-size:28px">{_fmt_pct(result.suggested_equity_weight)}</div>
        <div class="hero-label">Suggested Equity Weight (of NAV)</div>
        <div class="hero-label">Leverage: {_esc(result.leverage_note)}</div>
      </div>
    </div>
  </div>

  <div class="card">
    <div class="narrative">{_esc(result.narrative)}</div>
  </div>

  <div class="two-col">
    <div class="card">
      <h2>Group Scores</h2>
      {group_rows_html}
    </div>
    <div class="card">
      <h2>Polo Allocation Engine (% of NAV)</h2>
      <div class="alloc-bar">{alloc_segments}</div>
      <div class="legend">{alloc_legend}</div>
    </div>
  </div>

  <div class="card">
    <h2>Data freshness</h2>
    <div class="two-col">
      <div>
        <div style="font-size:12px;color:var(--text-muted);margin-bottom:6px;font-weight:600;">Auto-updated this run</div>
        <ul>{auto_list_html}</ul>
      </div>
      <div>
        <div style="font-size:12px;color:var(--text-muted);margin-bottom:6px;font-weight:600;">Needs manual weekly input</div>
        <ul>Valuation table, macro (GDP/credit/SBV/rates/CPI), margin &amp; ETF flow, sentiment/news/FDI — see the sheet's own Notes/Frequency columns for sources (Vietcap IQ, TCBS, VBMA, Fiintrade, Topi).</ul>
      </div>
    </div>
    {f'<div style="margin-top:12px;font-size:12px;color:var(--text-muted);font-weight:600;">Fetch issues this run</div><ul>{errors_html}</ul>' if errors_html else ''}
    {f'<div style="margin-top:12px;font-size:12px;color:var(--text-muted);font-weight:600;">Detected issues in the source workbook</div><ul class="warn-list">{warnings_html}</ul>' if warnings_html else ''}
  </div>

  <div class="card">
    <h2>All indicators</h2>
    <table>
      <thead><tr><th>Group</th><th>Factor</th><th class="num">Current</th><th class="num">Score /10</th><th>Source</th></tr></thead>
      <tbody>{indicator_rows_html}</tbody>
    </table>
  </div>
</div>
</body>
</html>"""


def write_dashboard(result: ScorecardResult, auto: AutoData, applied_fields: list, path: str) -> None:
    import os

    html_str = render_dashboard_html(result, auto, applied_fields)
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(html_str)
